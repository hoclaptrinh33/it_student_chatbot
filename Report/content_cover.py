# -*- coding: utf-8 -*-
"""
content_cover.py
Xây dựng Trang bìa chuẩn giống hệt bản báo cáo gốc:
- Trường ĐH Công nghệ Đông Á, Khoa CNTT
- BÀI TẬP LỚN, HỌC PHẦN: TRÍ TUỆ NHÂN TẠO, MÃ ĐỀ TÀI: 32
- TÊN ĐỀ TÀI: XÂY DỰNG CHƯƠNG TRÌNH HỎI ĐÁP BẰNG NGÔN NGỮ TỰ NHIÊN HỖ TRỢ CHỌN MÔN HỌC VÀ TÀI LIỆU HỌC TẬP CHO SINH VIÊN KHOA CNTT
- LỚP TÍN CHỈ: TTNT.01.K14.08.LH.C01_LT, NHÓM THỰC HIỆN: NHÓM 15
- Bảng danh sách 5 sinh viên ngay trên trang bìa
- BẮC NINH - 2026
Sau đó: Lời cảm ơn, Mục lục, Danh mục từ viết tắt.
"""

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from make_full_report import (
    FONT_NAME, BLACK, set_run_font, add_p, add_heading_1, add_heading_2, 
    add_bullet, add_styled_table
)

from docx.enum.table import WD_TABLE_ALIGNMENT

def build_cover_and_preamble(doc):
    # =========================================================================
    # TRANG BÌA CHÍNH (GIỐNG HỆT BẢN GỐC CỦA BÁO CÁO)
    # =========================================================================
    p0 = doc.add_paragraph()
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.paragraph_format.space_before = Pt(10)
    p0.paragraph_format.space_after = Pt(2)
    p0.paragraph_format.line_spacing = 1.5
    r0 = p0.add_run("TRƯỜNG ĐẠI HỌC CÔNG NGHỆ ĐÔNG Á")
    set_run_font(r0, FONT_NAME, 14, bold=True, color=BLACK)

    p1 = doc.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_before = Pt(0)
    p1.paragraph_format.space_after = Pt(25)
    p1.paragraph_format.line_spacing = 1.5
    r1 = p1.add_run("KHOA CÔNG NGHỆ THÔNG TIN")
    set_run_font(r1, FONT_NAME, 14, bold=True, color=BLACK)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(25)
    p2.paragraph_format.line_spacing = 1.5

    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_before = Pt(15)
    p3.paragraph_format.space_after = Pt(0)
    p3.paragraph_format.line_spacing = 1.5
    r3 = p3.add_run("BÀI TẬP LỚN")
    set_run_font(r3, FONT_NAME, 20, bold=True, color=BLACK)

    p4 = doc.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p4.paragraph_format.space_before = Pt(4)
    p4.paragraph_format.space_after = Pt(4)
    r4 = p4.add_run("HỌC PHẦN: TRÍ TUỆ NHÂN TẠO")
    set_run_font(r4, FONT_NAME, 14, bold=True, color=BLACK)

    p5 = doc.add_paragraph()
    p5.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p5.paragraph_format.space_before = Pt(4)
    p5.paragraph_format.space_after = Pt(4)
    r5 = p5.add_run("MÃ ĐỀ TÀI: 32")
    set_run_font(r5, FONT_NAME, 14, bold=True, color=BLACK)

    p6 = doc.add_paragraph()
    p6.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p6.paragraph_format.space_before = Pt(4)
    p6.paragraph_format.space_after = Pt(6)
    p6.paragraph_format.line_spacing = 1.2
    r6 = p6.add_run("TÊN ĐỀ TÀI: XÂY DỰNG CHƯƠNG TRÌNH HỎI ĐÁP BẰNG NGÔN NGỮ TỰ NHIÊN HỖ TRỢ CHỌN MÔN HỌC VÀ TÀI LIỆU HỌC TẬP CHO SINH VIÊN KHOA CNTT")
    set_run_font(r6, FONT_NAME, 14, bold=True, color=BLACK)

    p7 = doc.add_paragraph()
    p7.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p7.paragraph_format.space_before = Pt(4)
    p7.paragraph_format.space_after = Pt(12)
    p7.paragraph_format.line_spacing = 1.2
    r7 = p7.add_run("LỚP TÍN CHỈ: TTNT.01.K14.08.LH.C01_LT\nNHÓM THỰC HIỆN: NHÓM 15")
    set_run_font(r7, FONT_NAME, 14, bold=True, color=BLACK)

    # Bảng danh sách thành viên nhóm đặt ngay trên trang bìa (giống hệt bản gốc Table 0)
    tbl_cover = doc.add_table(rows=6, cols=4)
    tbl_cover.style = 'Table Grid'
    tbl_cover.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers_cover = ["STT", "Mã sinh viên", "Sinh viên thực hiện", "Lớp hành chính"]
    rows_cover = [
        ["1", "20233301", "Lê Hải Đăng", "DCCNTT14.9"],
        ["2", "20233302", "Lê Minh Quân", "DCCNTT14.9"],
        ["3", "20233303", "Lê Xuân Đạt", "DCCNTT14.9"],
        ["4", "20233304", "Lê Thanh Tùng", "DCCNTT14.9"],
        ["5", "20233305", "Phạm Bảo Sơn", "DCCNTT14.9"]
    ]
    col_widths = [Inches(0.8), Inches(1.8), Inches(2.4), Inches(1.6)]
    for i, h in enumerate(headers_cover):
        c = tbl_cover.cell(0, i)
        c.text = ""
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(h)
        set_run_font(r, FONT_NAME, 14, bold=True, color=BLACK)

    for r_idx, r_data in enumerate(rows_cover):
        for c_idx, val in enumerate(r_data):
            c = tbl_cover.cell(r_idx + 1, c_idx)
            c.text = ""
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(val)
            set_run_font(r, FONT_NAME, 14, bold=False, color=BLACK)

    for row in tbl_cover.rows:
        for c_idx, w in enumerate(col_widths):
            row.cells[c_idx].width = w

    p8 = doc.add_paragraph()
    p8.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p8.paragraph_format.space_before = Pt(40)
    p8.paragraph_format.space_after = Pt(0)
    p8.paragraph_format.line_spacing = 1.5
    r8 = p8.add_run("BẮC NINH - 2026")
    set_run_font(r8, FONT_NAME, 14, bold=True, color=BLACK)

    doc.add_page_break()

    # =========================================================================
    # LỜI CẢM ƠN
    # =========================================================================
    add_heading_1(doc, "LỜI CẢM ƠN")
    add_p(doc, "Để hoàn thành tốt bài tập lớn học phần Trí tuệ nhân tạo với đề tài “Xây dựng chương trình hỏi đáp bằng ngôn ngữ tự nhiên hỗ trợ chọn môn học và tài liệu học tập cho sinh viên khoa CNTT”, nhóm sinh viên chúng em xin được bày tỏ lòng biết ơn sâu sắc và chân thành nhất tới các thầy, cô giáo trong Khoa Công nghệ Thông tin - Trường Đại học Công nghệ Đông Á.")
    add_p(doc, "Đặc biệt, nhóm chúng em xin gửi lời cảm ơn trân trọng tới ThS. Nguyễn Văn An, giảng viên trực tiếp phụ trách giảng dạy và hướng dẫn học phần Trí tuệ nhân tạo. Trong suốt thời gian thực hiện nghiên cứu, thầy đã tận tình truyền đạt những kiến thức chuyên môn quý báu về Tác tử thông minh, Không gian trạng thái, Thuật toán tìm kiếm, Lập luận xác suất Bayes, Học máy và các kiến trúc AI hiện đại. Thầy cũng đã đưa ra những định hướng khoa học, góp ý sâu sắc và tạo mọi điều kiện thuận lợi nhất để nhóm có thể hiện thực hóa thành công sản phẩm từ lý thuyết sang ứng dụng thực tiễn.")
    add_p(doc, "Mặc dù nhóm đã nỗ lực hết mình, chủ động nghiên cứu và áp dụng các công nghệ tiên tiến nhất như FastAPI, PostgreSQL 15, Qdrant Vector Database, Vietnamese-SBERT và Mô hình ngôn ngữ lớn (LLM), song do phạm vi rộng của đề tài và giới hạn về mặt thời gian, bài báo cáo khó tránh khỏi những thiếu sót nhất định. Nhóm chúng em rất mong nhận được những nhận xét, đóng góp quý báu từ quý thầy cô và các bạn sinh viên để sản phẩm ngày càng hoàn thiện, có tính ứng dụng cao hơn nữa trong thực tiễn đào tạo của Nhà trường.")
    add_p(doc, "Chúng em xin chân thành cảm ơn!")

    p_sign = doc.add_paragraph()
    p_sign.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sign.paragraph_format.space_before = Pt(20)
    p_sign.paragraph_format.space_after = Pt(0)
    r_sign = p_sign.add_run("Bắc Ninh, tháng 09 năm 2026\nĐại diện Nhóm thực hiện\n\n\nLê Hải Đăng")
    set_run_font(r_sign, FONT_NAME, 13, bold=True, italic=True, color=BLACK)

    doc.add_page_break()

    # =========================================================================
    # MỤC LỤC
    # =========================================================================
    add_heading_1(doc, "MỤC LỤC")
    toc_items = [
        ("PHẦN MỞ ĐẦU", "1"),
        ("1. Lý do chọn đề tài", "1"),
        ("2. Mục tiêu của đề tài", "2"),
        ("3. Đối tượng và phạm vi nghiên cứu", "3"),
        ("4. Phương pháp nghiên cứu", "3"),
        ("5. Cấu trúc báo cáo", "4"),
        ("CHƯƠNG 1: CƠ SỞ LÝ THUYẾT VÀ CÁC PHƯƠNG PHÁP TRÍ TUỆ NHÂN TẠO", "5"),
        ("1.1 Tổng quan về Trí tuệ nhân tạo và Tác tử thông minh trong Hệ thống Hỏi - Đáp", "5"),
        ("1.2 Không gian trạng thái và Các phương pháp tìm kiếm trong Quản lý Lộ trình Học tập", "9"),
        ("1.3 Lập luận xác suất và Xử lý thông tin không chắc chắn", "13"),
        ("1.4 Tổng quan Học máy và Xử lý ngôn ngữ tự nhiên trong Kiến trúc Chatbot RAG", "16"),
        ("1.5 Phương pháp đánh giá mô hình và hệ thống Hỏi - Đáp", "28"),
        ("1.6 Công cụ, Thư viện và Hạ tầng công nghệ phát triển hệ thống", "31"),
        ("CHƯƠNG 2: PHÂN TÍCH VÀ XÂY DỰNG CHƯƠNG TRÌNH", "33"),
        ("2.1 Phát biểu bài toán", "33"),
        ("2.2 Xác định yêu cầu, Input và Output của hệ thống", "35"),
        ("2.3 Thiết kế sơ đồ khối và kiến trúc hệ thống", "38"),
        ("2.4 Mô tả thuật toán và giải thuật xử lý cốt lõi", "44"),
        ("2.5 Mô tả dữ liệu (Dataset) và Cơ sở tri thức", "50"),
        ("2.6 Cài đặt hệ thống và Hệ thống sơ đồ kỹ thuật & Giao diện tương tác", "55"),
        ("CHƯƠNG 3: THỰC NGHIỆM VÀ ĐÁNH GIÁ", "72"),
        ("3.1 Môi trường thực nghiệm và triển khai", "72"),
        ("3.2 Kịch bản kiểm thử chương trình với các trường hợp thực tế", "73"),
        ("3.3 Đánh giá kết quả thực nghiệm", "78"),
        ("3.4 Đánh giá ưu điểm và hạn chế của hệ thống", "81"),
        ("PHẦN KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN", "83"),
        ("TÀI LIỆU THAM KHẢO", "85"),
        ("PHỤ LỤC: LIÊN KẾT MÃ NGUỒN GITHUB VÀ HƯỚNG DẪN DEMO", "87")
    ]
    for title, page in toc_items:
        p_toc = doc.add_paragraph()
        p_toc.paragraph_format.space_before = Pt(1)
        p_toc.paragraph_format.space_after = Pt(2)
        p_toc.paragraph_format.line_spacing = 1.15
        is_bold = title.startswith("CHƯƠNG") or title.startswith("PHẦN") or title.startswith("TÀI")
        r_t = p_toc.add_run(title)
        set_run_font(r_t, FONT_NAME, 12, bold=is_bold, color=BLACK)
        
        dots_len = max(5, 75 - len(title))
        r_dots = p_toc.add_run(" " + "." * dots_len + " ")
        set_run_font(r_dots, FONT_NAME, 11, bold=False, color=BLACK)
        
        r_p = p_toc.add_run(page)
        set_run_font(r_p, FONT_NAME, 12, bold=is_bold, color=BLACK)

    doc.add_page_break()

    # =========================================================================
    # DANH MỤC CÁC CHỮ VIẾT TẮT
    # =========================================================================
    add_heading_1(doc, "DANH MỤC CÁC CHỮ VIẾT TẮT")
    abbr_headers = ["Ký hiệu viết tắt", "Thuật ngữ tiếng Anh", "Ý nghĩa tiếng Việt"]
    abbr_rows = [
        ["AI", "Artificial Intelligence", "Trí tuệ nhân tạo"],
        ["RAG", "Retrieval-Augmented Generation", "Tạo sinh tăng cường truy xuất"],
        ["LLM", "Large Language Model", "Mô hình ngôn ngữ lớn"],
        ["NLP", "Natural Language Processing", "Xử lý ngôn ngữ tự nhiên"],
        ["NLU", "Natural Language Understanding", "Hiểu ngôn ngữ tự nhiên"],
        ["ML", "Machine Learning", "Học máy"],
        ["ERD", "Entity-Relationship Diagram", "Sơ đồ thực thể liên kết cơ sở dữ liệu"],
        ["DAG", "Directed Acyclic Graph", "Đồ thị có hướng không chu trình"],
        ["BFS", "Breadth-First Search", "Thuật toán tìm kiếm theo chiều rộng"],
        ["DFS", "Depth-First Search", "Thuật toán tìm kiếm theo chiều sâu"],
        ["RBAC", "Role-Based Access Control", "Kiểm soát truy cập dựa trên vai trò"],
        ["API", "Application Programming Interface", "Giao diện lập trình ứng dụng"],
        ["JWT", "JSON Web Token", "Mã thông báo định danh Web JSON"],
        ["SBERT", "Sentence-BERT", "Mô hình nhúng câu dựa trên kiến trúc BERT"],
        ["TF-IDF", "Term Frequency - Inverse Document Frequency", "Tần số xuất hiện từ - Nghịch đảo tần số tài liệu"],
        ["ANN", "Approximate Nearest Neighbor", "Thuật toán tìm kiếm lân cận gần đúng"],
        ["HNSW", "Hierarchical Navigable Small World", "Đồ thị thế giới nhỏ phân cấp điều hướng"],
        ["MRR", "Mean Reciprocal Rank", "Điểm trung bình nghịch đảo thứ hạng"],
        ["MSE", "Mean Squared Error", "Sai số toàn phương trung bình"],
        ["MAE", "Mean Absolute Error", "Sai số tuyệt đối trung bình"],
        ["ROC-AUC", "Receiver Operating Characteristic - Area Under Curve", "Diện tích dưới đường cong ROC"],
        ["MLOps", "Machine Learning Operations", "Vận hành và quản lý vòng đời học máy"]
    ]
    add_styled_table(doc, abbr_headers, abbr_rows, [1.2, 2.5, 2.7], "Bảng 0.2: Danh mục các thuật ngữ và chữ viết tắt trong báo cáo")

    doc.add_page_break()
